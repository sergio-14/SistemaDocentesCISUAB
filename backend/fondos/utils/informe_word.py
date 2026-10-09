"""Informe de Fondo de Tiempo en Word (.docx), con el mismo contenido y forma que
el PDF (InformePDFGenerator): membrete con logo en cada página, fecha, A/DE/REF,
saludo, las secciones del docente, cierre, firma (con su imagen) y número de
página. El HTML del editor se pasa a Word con las mismas etiquetas que entiende
el PDF: negrita, cursiva, subrayado, tachado, tamaño y color de letra,
alineación, interlineado, sangría, listas, tablas e imágenes.
"""
import base64
import io
import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from fondos.models import SECCIONES_INFORME
from fondos.utils.informe_imagenes import leer_imagen as leer_imagen_informe
from fondos.utils.informe_texto import CAMPOS_TEXTO_INFORME, construir_defaults_informe
from fondos.utils.pdf_generator import _normalizar_color_css, _parsear_estilo_css, _tamano_fuente_a_pt

FUENTE = 'Verdana'
FUENTE_FIRMA = 'Trebuchet MS'
TAMANO_TEXTO = Pt(10)
COLOR_LINEA_MEMBRETE = '3CBECA'
ANCHO_PAGINA = Cm(21.59)
ALTO_PAGINA = Cm(27.94)
MARGEN_LATERAL = Cm(2.5)
ANCHO_CONTENIDO = ANCHO_PAGINA - 2 * MARGEN_LATERAL
ALTO_MAXIMO_IMAGEN = Cm(9)

_ALINEACION = {
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_DATA_IMG_RE = re.compile(r'^data:image/(png|jpe?g|gif);base64,(?P<data>.+)$', re.IGNORECASE | re.DOTALL)
_IMG_SRC_RE = re.compile(r'<img\b[^>]*?\bsrc\s*=\s*(["\'])(.*?)\1', re.IGNORECASE | re.DOTALL)


def _borde_inferior(parrafo, color, grosor=8):
    """Línea bajo el párrafo (la franja turquesa del membrete)."""
    pPr = parrafo._p.get_or_add_pPr()
    bordes = OxmlElement('w:pBdr')
    inferior = OxmlElement('w:bottom')
    for clave, valor in (('w:val', 'single'), ('w:sz', str(grosor)), ('w:space', '1'), ('w:color', color)):
        inferior.set(qn(clave), valor)
    bordes.append(inferior)
    pPr.append(bordes)


def _campo_pagina(parrafo):
    """Número de página automático (campo PAGE de Word)."""
    run = parrafo.add_run()
    for tipo, texto in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tipo:
            elemento = OxmlElement('w:fldChar')
            elemento.set(qn('w:fldCharType'), tipo)
        else:
            elemento = OxmlElement('w:instrText')
            elemento.set(qn('xml:space'), 'preserve')
            elemento.text = texto
        run._r.append(elemento)
    run.font.size = Pt(9)


def _formato_word_actual(documento):
    """Deja el documento como uno de Word actual: lo marca como de Word 2013 o
    posterior (la plantilla de python-docx no lo declara y Word lo abre en "Modo de
    compatibilidad") y le quita la miniatura de la plantilla."""
    settings = documento.settings.element
    compat = settings.find(qn('w:compat'))
    if compat is None:
        compat = OxmlElement('w:compat')
        settings.append(compat)
    for ajuste in compat.findall(qn('w:compatSetting')):
        if ajuste.get(qn('w:name')) == 'compatibilityMode':
            compat.remove(ajuste)
    ajuste = OxmlElement('w:compatSetting')
    for clave, valor in (
        ('w:name', 'compatibilityMode'), ('w:uri', 'http://schemas.microsoft.com/office/word'), ('w:val', '15'),
    ):
        ajuste.set(qn(clave), valor)
    compat.append(ajuste)

    # La plantilla trae una miniatura de una hoja en blanco: Windows la usaría como
    # ícono del archivo en vez del ícono de Word.
    relaciones = documento.part.package.rels
    for rId, relacion in list(relaciones.items()):
        if relacion.reltype == RT.THUMBNAIL:
            relaciones.pop(rId)


def _sin_bordes(tabla):
    """Tabla de maquetación: sin líneas."""
    tblPr = tabla._tbl.tblPr
    bordes = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        elemento = OxmlElement(f'w:{lado}')
        elemento.set(qn('w:val'), 'nil')
        bordes.append(elemento)
    tblPr.append(bordes)


def _ancho_columna(tabla, indice, ancho):
    for celda in tabla.columns[indice].cells:
        celda.width = ancho


def _parrafo(contenedor, texto='', *, negrita=False, cursiva=False, tamano=None, alineacion=None,
             fuente=None, antes=0, despues=4, color=None, subrayado=False):
    parrafo = contenedor.add_paragraph()
    formato = parrafo.paragraph_format
    formato.space_before = Pt(antes)
    formato.space_after = Pt(despues)
    if alineacion is not None:
        parrafo.alignment = alineacion
    if texto:
        run = parrafo.add_run(texto)
        run.bold = negrita
        run.italic = cursiva
        run.underline = subrayado
        if tamano:
            run.font.size = Pt(tamano)
        if fuente:
            run.font.name = fuente
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return parrafo


def _tamano_imagen(img_bytes, ancho_css, ancho_maximo, alto_maximo):
    """(ancho, alto) en EMU de una imagen: el ancho del editor si lo trae, sin pasar
    del espacio disponible y sin deformarla. None si no se puede leer."""
    from PIL import Image as ImagenPIL
    try:
        with ImagenPIL.open(io.BytesIO(img_bytes)) as imagen:
            ancho_px, alto_px = imagen.size
    except Exception:
        return None
    if not ancho_px or not alto_px:
        return None
    ancho = Pt(ancho_px * 0.75)
    if ancho_css:
        ancho_css = ancho_css.strip()
        if ancho_css.endswith('%'):
            try:
                ancho = int(ancho_maximo * float(ancho_css[:-1]) / 100)
            except ValueError:
                pass
        else:
            puntos = _tamano_fuente_a_pt(ancho_css, None)
            if puntos:
                ancho = Pt(puntos)
    ancho = min(ancho, ancho_maximo)
    alto = int(ancho * alto_px / ancho_px)
    if alto > alto_maximo:
        ancho = int(ancho * alto_maximo / alto)
        alto = alto_maximo
    return int(ancho), int(alto)


class _HtmlAWord(HTMLParser):
    """Escribe en `contenedor` (el documento o una celda) el HTML del editor del informe."""

    _CORTES = {'p', 'div', 'li', 'h1', 'h2', 'h3', 'h4'}

    def __init__(self, contenedor, ancho_maximo=ANCHO_CONTENIDO):
        super().__init__(convert_charrefs=True)
        self.contenedor = contenedor
        self.ancho_maximo = ancho_maximo
        self.parrafo = None
        self.formatos = []  # pila de dicts de formato de texto
        self.alineacion = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.interlineado = None
        self.sangria = 0
        self.listas = []  # [[tipo, contador]]
        self.prefijo_pendiente = None
        self.tabla = None  # [[(texto, es_encabezado)]]
        self.celda = None

    # --- formato de texto ---
    def _formato_actual(self):
        actual = {}
        for formato in self.formatos:
            actual.update(formato)
        return actual

    def _nuevo_parrafo(self):
        self.parrafo = self.contenedor.add_paragraph()
        formato = self.parrafo.paragraph_format
        formato.space_after = Pt(4)
        formato.left_indent = Cm(self.sangria)
        if self.interlineado:
            formato.line_spacing = self.interlineado
        self.parrafo.alignment = self.alineacion
        if self.prefijo_pendiente:
            self.parrafo.add_run(self.prefijo_pendiente)
            self.prefijo_pendiente = None
        return self.parrafo

    def _cerrar_parrafo(self):
        self.parrafo = None
        self.alineacion = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.interlineado = None

    def _estilo_bloque(self, attrs, estilo):
        alineacion = estilo.get('text-align') or attrs.get('align')
        if alineacion in _ALINEACION:
            self.alineacion = _ALINEACION[alineacion]
        if estilo.get('line-height'):
            try:
                self.interlineado = float(estilo['line-height'])
            except ValueError:
                pass

    def _formato_de_etiqueta(self, tag, attrs, estilo):
        formato = {}
        if tag in ('b', 'strong') or estilo.get('font-weight') in ('bold', '700', '600'):
            formato['bold'] = True
        if tag in ('i', 'em') or estilo.get('font-style') == 'italic':
            formato['italic'] = True
        decoracion = estilo.get('text-decoration', '') + ' ' + estilo.get('text-decoration-line', '')
        if tag == 'u' or 'underline' in decoracion:
            formato['underline'] = True
        if tag in ('s', 'strike', 'del') or 'line-through' in decoracion:
            formato['strike'] = True
        tamano = _tamano_fuente_a_pt(estilo.get('font-size'), attrs.get('size') if tag == 'font' else None)
        if tamano:
            formato['size'] = tamano
        color = _normalizar_color_css(estilo.get('color') or (attrs.get('color') if tag == 'font' else None))
        if color and re.fullmatch(r'#[0-9a-f]{6}', color):
            formato['color'] = color[1:]
        return formato

    # --- HTMLParser ---
    def handle_starttag(self, tag, attrs_lista):
        attrs = dict(attrs_lista)
        estilo = _parsear_estilo_css(attrs.get('style', ''))

        if self.tabla is not None:
            if tag == 'tr':
                self.tabla.append([])
            elif tag in ('td', 'th'):
                self.celda = [[], tag == 'th']
            elif tag == 'br' and self.celda is not None:
                self.celda[0].append('\n')
            return
        if tag == 'table':
            self._cerrar_parrafo()
            self.tabla = []
            return

        if tag in self._CORTES:
            self._cerrar_parrafo()
            self._estilo_bloque(attrs, estilo)
            if tag == 'li':
                tipo_contador = self.listas[-1] if self.listas else ['ul', 0]
                if tipo_contador[0] == 'ol':
                    tipo_contador[1] += 1
                    self.prefijo_pendiente = f'{tipo_contador[1]}.  '
                else:
                    self.prefijo_pendiente = '•  '
            if tag.startswith('h'):
                self.formatos.append({'bold': True})
            return
        if tag in ('ul', 'ol'):
            self._cerrar_parrafo()
            self.listas.append([tag, 0])
            self.sangria += 0.6
            return
        if tag == 'blockquote':
            self._cerrar_parrafo()
            self.sangria += 1
            return
        if tag == 'br':
            (self.parrafo or self._nuevo_parrafo()).add_run().add_break(WD_BREAK.LINE)
            return
        if tag == 'img':
            self._imagen(attrs, estilo)
            return
        if tag in ('b', 'strong', 'i', 'em', 'u', 's', 'strike', 'del', 'span', 'font'):
            self.formatos.append(self._formato_de_etiqueta(tag, attrs, estilo))

    def handle_endtag(self, tag):
        if self.tabla is not None:
            if tag in ('td', 'th') and self.celda is not None and self.tabla:
                texto = re.sub(r'[ \t\r\f\v]+', ' ', ''.join(self.celda[0])).strip()
                self.tabla[-1].append((texto, self.celda[1]))
                self.celda = None
            elif tag == 'table':
                self._escribir_tabla()
            return
        if tag in self._CORTES:
            if tag.startswith('h') and self.formatos:
                self.formatos.pop()
            self._cerrar_parrafo()
        elif tag in ('ul', 'ol'):
            if self.listas:
                self.listas.pop()
            self.sangria = max(0, self.sangria - 0.6)
            self._cerrar_parrafo()
        elif tag == 'blockquote':
            self.sangria = max(0, self.sangria - 1)
            self._cerrar_parrafo()
        elif tag in ('b', 'strong', 'i', 'em', 'u', 's', 'strike', 'del', 'span', 'font') and self.formatos:
            self.formatos.pop()

    def handle_data(self, data):
        if self.tabla is not None:
            if self.celda is not None:
                self.celda[0].append(data)
            return
        texto = re.sub(r'\s+', ' ', data)
        if not texto.strip() and self.parrafo is None:
            return
        parrafo = self.parrafo or self._nuevo_parrafo()
        # Sin espacios al inicio del párrafo (ni tras la viñeta o el número de la lista).
        if re.fullmatch(r'\s*(•|\d+\.)?\s*', parrafo.text):
            texto = texto.lstrip()
        if not texto:
            return
        run = parrafo.add_run(texto)
        formato = self._formato_actual()
        run.bold = formato.get('bold')
        run.italic = formato.get('italic')
        run.underline = formato.get('underline')
        run.font.strike = formato.get('strike')
        if formato.get('size'):
            run.font.size = Pt(formato['size'])
        if formato.get('color'):
            run.font.color.rgb = RGBColor.from_string(formato['color'].upper())

    # --- bloques especiales ---
    def _imagen(self, attrs, estilo):
        img_bytes = leer_imagen_informe(attrs.get('src', '') or '')
        if not img_bytes:
            return
        medidas = _tamano_imagen(img_bytes, estilo.get('width'), self.ancho_maximo, ALTO_MAXIMO_IMAGEN)
        if not medidas:
            return
        alineacion = self.alineacion
        self._cerrar_parrafo()
        parrafo = self.contenedor.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT if alineacion == WD_ALIGN_PARAGRAPH.JUSTIFY else alineacion
        parrafo.paragraph_format.space_before = Pt(4)
        parrafo.paragraph_format.space_after = Pt(4)
        try:
            parrafo.add_run().add_picture(io.BytesIO(img_bytes), width=medidas[0], height=medidas[1])
        except Exception:
            pass
        self.alineacion = alineacion

    def _escribir_tabla(self):
        filas = [fila for fila in self.tabla if fila]
        self.tabla = None
        if not filas:
            return
        columnas = max(len(fila) for fila in filas)
        tabla = self.contenedor.add_table(rows=len(filas), cols=columnas)
        tabla.style = 'Table Grid'
        for i, fila in enumerate(filas):
            for j, (texto, encabezado) in enumerate(fila):
                celda = tabla.cell(i, j)
                celda.text = ''
                run = celda.paragraphs[0].add_run(texto)
                run.bold = encabezado
                run.font.size = Pt(9)
        self.contenedor.add_paragraph().paragraph_format.space_after = Pt(2)

    def finalizar(self):
        self.close()
        if self.tabla is not None:
            self._escribir_tabla()


def _escribir_html(contenedor, html):
    parser = _HtmlAWord(contenedor)
    parser.feed(html or '')
    parser.finalizar()


def _logo_carrera(fondo):
    if not fondo.carrera:
        return None
    try:
        match = _DATA_IMG_RE.match(fondo.carrera.get_logo_carrera_data_uri() or '')
        return base64.b64decode(match.group('data')) if match else None
    except Exception:
        return None


def _escribir_membrete(contenedor, lineas, logo):
    """Logo y líneas del encabezado institucional, con la franja turquesa debajo."""
    tabla = contenedor.add_table(rows=1, cols=3, width=ANCHO_CONTENIDO)
    _sin_bordes(tabla)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for indice, ancho in enumerate((Cm(2.2), ANCHO_CONTENIDO - Cm(4.4), Cm(2.2))):
        _ancho_columna(tabla, indice, ancho)
    if logo:
        try:
            tabla.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(logo), width=Cm(1.8))
        except Exception:
            pass
    celda_texto = tabla.cell(0, 1)
    for indice, linea in enumerate(lineas):
        parrafo = celda_texto.paragraphs[0] if indice == 0 else celda_texto.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo.paragraph_format.space_after = Pt(0)
        run = parrafo.add_run(linea)
        run.bold = True
        run.font.size = Pt(9.5 if indice == 0 else 8.5)
    franja = contenedor.add_paragraph()
    franja.paragraph_format.space_after = Pt(14)
    _borde_inferior(franja, COLOR_LINEA_MEMBRETE, grosor=12)


def _membrete(seccion, fondo, encabezado_texto):
    """Membrete como encabezado de página, igual en todas las hojas (como en el PDF):
    es uno solo, así que al editarlo en Word cambia en todas. Word lo muestra más
    claro en pantalla mientras se edita el cuerpo, pero lo imprime normal. El número
    de página va en el pie."""
    lineas = [linea.strip() for linea in (encabezado_texto or '').split('\n') if linea.strip()][:4]
    if not lineas:
        lineas = [fondo.carrera.nombre.upper() if fondo.carrera else 'CARRERA']

    cabecera = seccion.header
    cabecera.is_linked_to_previous = False
    base = cabecera.paragraphs[0]
    _escribir_membrete(cabecera, lineas, _logo_carrera(fondo))
    base._p.getparent().remove(base._p)

    pie = seccion.footer
    pie.is_linked_to_previous = False
    numero = pie.paragraphs[0]
    numero.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _campo_pagina(numero)


def _firma(documento, textos, firma_imagen_html):
    """Bloque de la firma a la derecha: imagen (si hay), nombre, cargo y correo centrados entre sí."""
    tabla = documento.add_table(rows=1, cols=1)
    _sin_bordes(tabla)
    tabla.alignment = WD_TABLE_ALIGNMENT.RIGHT
    _ancho_columna(tabla, 0, Cm(8.5))
    celda = tabla.cell(0, 0)
    primero = celda.paragraphs[0]
    primero.alignment = WD_ALIGN_PARAGRAPH.CENTER

    match = _IMG_SRC_RE.search(firma_imagen_html or '')
    img_bytes = leer_imagen_informe(match.group(2)) if match else None
    medidas = _tamano_imagen(img_bytes, None, Cm(6), Cm(2.5)) if img_bytes else None
    if medidas:
        try:
            primero.add_run().add_picture(io.BytesIO(img_bytes), width=medidas[0], height=medidas[1])
            primero = celda.add_paragraph()
            primero.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass

    run = primero.add_run(textos['firma_nombre'])
    run.italic = True
    run.font.name = FUENTE_FIRMA
    run.font.size = Pt(13)
    primero.paragraph_format.space_after = Pt(0)
    _parrafo(celda, textos['firma_cargo'], negrita=True, tamano=9, alineacion=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
    if textos['firma_email']:
        parrafo = _parrafo(celda, 'E-mail: ', tamano=7, alineacion=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
        correo = parrafo.add_run(textos['firma_email'])
        correo.underline = True
        correo.font.size = Pt(7)
        correo.font.color.rgb = RGBColor.from_string('0563C1')


def generar_informe_word(fondo, informe):
    """BytesIO con el .docx del informe (guardado o, si aún no hay, con los textos precargados)."""
    defaults = construir_defaults_informe(fondo)
    textos = {}
    for campo in CAMPOS_TEXTO_INFORME:
        guardado = (getattr(informe, campo, '') or '').strip() if informe else ''
        textos[campo] = guardado or defaults[campo]

    documento = Document()
    _formato_word_actual(documento)
    estilo = documento.styles['Normal']
    estilo.font.name = FUENTE
    estilo.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
    estilo.font.size = TAMANO_TEXTO
    estilo.paragraph_format.line_spacing = 1.2

    seccion = documento.sections[0]
    seccion.start_type = WD_SECTION.NEW_PAGE
    seccion.page_width, seccion.page_height = ANCHO_PAGINA, ALTO_PAGINA
    seccion.left_margin = seccion.right_margin = MARGEN_LATERAL
    seccion.top_margin = Cm(3.8)
    seccion.bottom_margin = Cm(1.8)
    seccion.header_distance = Cm(0.8)
    seccion.footer_distance = Cm(0.8)
    _membrete(seccion, fondo, textos['encabezado_texto'])

    # 1. Fecha
    _parrafo(documento, textos['fecha_texto'], alineacion=WD_ALIGN_PARAGRAPH.RIGHT, despues=18)

    # 2. A / DE / REF
    filas = [
        ('A', textos['destinatario_nombre'], textos['destinatario_cargo'], False),
        ('DE', textos['remitente_nombre'], textos['remitente_cargo'], False),
        ('REF.', textos['referencia_texto'], None, True),
    ]
    tabla = documento.add_table(rows=len(filas), cols=2)
    _sin_bordes(tabla)
    _ancho_columna(tabla, 0, Cm(2.3))
    _ancho_columna(tabla, 1, ANCHO_CONTENIDO - Cm(2.3))
    for indice, (etiqueta, linea1, linea2, es_referencia) in enumerate(filas):
        tabla.cell(indice, 0).paragraphs[0].add_run(f'{etiqueta}\t:')
        valor = tabla.cell(indice, 1).paragraphs[0]
        run = valor.add_run(linea1)
        if es_referencia:
            run.bold = True
            run.underline = True
        if linea2:
            valor.add_run().add_break(WD_BREAK.LINE)
            valor.add_run(linea2).bold = True
    _parrafo(documento, despues=6)

    # 3. Saludo e introducción
    _escribir_html(documento, textos['saludo_intro_html'])

    # 4. Secciones (las opcionales vacías se omiten, como en el PDF)
    definicion = informe.secciones() if informe else SECCIONES_INFORME
    for numero, (campo, titulo, obligatoria) in enumerate(definicion, start=1):
        contenido = (getattr(informe, campo, '') or '').strip() if informe else ''
        if not contenido and not obligatoria:
            continue
        titulo_parrafo = _parrafo(documento, f'{numero}. {titulo.upper()}', negrita=True, antes=12, despues=4)
        titulo_parrafo.paragraph_format.keep_with_next = True
        if contenido:
            _escribir_html(documento, contenido)
        else:
            _parrafo(documento, 'No aplica / sin actividad en esta categoría.')

    # 5. Cierre, firma y copia
    _parrafo(documento, despues=6)
    _escribir_html(documento, textos['cierre_html'])
    _parrafo(documento, despues=30)
    _firma(documento, textos, getattr(informe, 'firma_imagen_html', '') if informe else '')
    _parrafo(documento, despues=0, antes=24)
    _parrafo(documento, 'Cc./Arch.')

    buffer = io.BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer
