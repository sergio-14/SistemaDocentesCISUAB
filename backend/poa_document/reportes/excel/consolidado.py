from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from .estilos import finalizar_tabla, guardar_en_buffer


def generar_consolidado_requerimientos_excel(items, gestion, total=None):
    items = list(items or [])
    total = Decimal(total if total is not None else sum(
        (Decimal(item['monto_estimado_total']) for item in items),
        Decimal('0'),
    ))
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Consolidado compras'
    hoja.append([f'Consolidado de requerimientos POA - Gestión {gestion}'])
    hoja.append(['Partida', 'Tipo', 'Ítem', 'Unidad', 'Características', 'Cantidad total', 'Costo estimado', 'Actividades'])
    for item in items:
        hoja.append([
            item['partida'], item['tipo'], item['item'], item['unidad_medida'],
            item['caracteristicas'], item['cantidad_total'],
            float(item['monto_estimado_total']), item['actividades_solicitantes'],
        ])
    hoja.append(['', '', '', '', 'TOTAL', '', float(total), ''])
    hoja.cell(hoja.max_row, 5).font = Font(bold=True)
    hoja.cell(hoja.max_row, 7).font = Font(bold=True)
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    hoja['A1'].font = Font(bold=True, size=14)
    finalizar_tabla(hoja, fila_encabezado=2)
    hoja.column_dimensions['C'].width = max(hoja.column_dimensions['C'].width, 38)
    hoja.column_dimensions['E'].width = max(hoja.column_dimensions['E'].width, 28)
    return guardar_en_buffer(libro)
