from io import BytesIO

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


AZUL_ENCABEZADO = '1D4ED8'
BORDE_SUAVE = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)


def aplicar_encabezado(hoja, fila=1, color=AZUL_ENCABEZADO):
    for celda in hoja[fila]:
        celda.font = Font(color='FFFFFF', bold=True)
        celda.fill = PatternFill(fill_type='solid', fgColor=color)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celda.border = BORDE_SUAVE


def aplicar_cuerpo(hoja, fila_inicio=2):
    for fila in hoja.iter_rows(min_row=fila_inicio):
        for celda in fila:
            celda.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            celda.border = BORDE_SUAVE


def ajustar_anchos(hoja, minimo=12, maximo=42):
    for indice in range(1, hoja.max_column + 1):
        largo = max((len(str(celda.value or '')) for celda in hoja[get_column_letter(indice)]), default=0)
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(largo + 2, minimo), maximo)


def finalizar_tabla(hoja, fila_encabezado=1):
    hoja.freeze_panes = f'A{fila_encabezado + 1}'
    hoja.auto_filter.ref = f'A{fila_encabezado}:{get_column_letter(hoja.max_column)}{hoja.max_row}'
    aplicar_encabezado(hoja, fila_encabezado)
    aplicar_cuerpo(hoja, fila_encabezado + 1)
    ajustar_anchos(hoja)


def guardar_en_buffer(libro):
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer
