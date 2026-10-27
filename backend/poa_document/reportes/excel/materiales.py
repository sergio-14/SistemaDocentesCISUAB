from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .estilos import guardar_en_buffer


def generar_informe_materiales_excel(datos, gestion):
    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    resumen.append([f'Seguimiento de abastecimiento POA · Gestión {gestion}'])
    resumen.append(['Indicador', 'Monto / cantidad'])
    for indicador, valor in datos['resumen']:
        resumen.append([indicador, valor])
    resumen.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

    hojas = {
        'Compras': ('compras', ['Proveedor', 'N.º compra', 'Fecha', 'Estado', 'Partida', 'Material', 'Unidad', 'Planificado', 'Comprado', 'Recibido', 'Por recibir', 'Costo unitario', 'Costo compra', 'Respaldo', 'Actividades de origen']),
        'Recepciones': ('recepciones', ['Proveedor', 'N.º compra', 'N.º acta/factura', 'Fecha', 'Estado', 'Partida', 'Material', 'Unidad', 'Cantidad recibida', 'Costo unitario', 'Total recibido', 'Recibido por', 'Respaldo']),
        'Stock': ('stock', ['Partida', 'Material', 'Unidad', 'Proveedor', 'N.º compra', 'Fecha recepción', 'Recibido', 'Entregado', 'Disponible', 'Costo unitario', 'Valor disponible', 'Estado']),
        'Entregas': ('entregas', ['Fecha', 'Proveedor', 'N.º compra', 'Partida', 'Material', 'Actividad', 'Unidad', 'Cantidad', 'Receptor', 'Documento', 'Cargo', 'Estado', 'Motivo de anulación', 'Acta']),
    }
    for nombre, (clave, encabezados) in hojas.items():
        hoja = libro.create_sheet(nombre)
        hoja.append(encabezados)
        for fila in datos[clave]:
            hoja.append(fila)

    for hoja in libro.worksheets:
        hoja.freeze_panes = 'A3' if hoja.title == 'Resumen' else 'A2'
        hoja.auto_filter.ref = hoja.dimensions if hoja.title != 'Resumen' else 'A2:B2'
        for celda in hoja[1]:
            celda.font = Font(name='Calibri', size=13 if hoja.title == 'Resumen' else 11, bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='17365D')
            celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if hoja.title == 'Resumen':
            for celda in hoja[2]:
                celda.font = Font(bold=True, color='FFFFFF')
                celda.fill = PatternFill('solid', fgColor='4472C4')
                celda.alignment = Alignment(horizontal='center', vertical='center')
            hoja.column_dimensions['A'].width = 38
            hoja.column_dimensions['B'].width = 24
        else:
            hoja.row_dimensions[1].height = 32
            for columna in hoja.columns:
                letra = columna[0].column_letter
                maximo = max((len(str(celda.value or '')) for celda in columna), default=10)
                hoja.column_dimensions[letra].width = min(max(maximo + 2, 12), 34)
                for celda in columna[1:]:
                    celda.alignment = Alignment(vertical='top', wrap_text=True)
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                if isinstance(celda.value, (int, float)):
                    celda.alignment = Alignment(horizontal='right', vertical='top')
        columnas_monetarias = {
            'Compras': ('L', 'M'),
            'Recepciones': ('J', 'K'),
            'Stock': ('J', 'K'),
        }.get(hoja.title, ())
        for letra in columnas_monetarias:
            for fila in range(2, hoja.max_row + 1):
                hoja[f'{letra}{fila}'].number_format = '#,##0.00'
        if hoja.title == 'Resumen':
            for fila in range(3, hoja.max_row + 1):
                if any(clave in str(hoja.cell(fila, 1).value).lower() for clave in ('presupuesto', 'costo', 'valor')):
                    hoja.cell(fila, 2).number_format = '#,##0.00'
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                if isinstance(celda.value, str) and celda.value.startswith('http'):
                    celda.hyperlink = celda.value
                    celda.style = 'Hyperlink'
    return guardar_en_buffer(libro)
