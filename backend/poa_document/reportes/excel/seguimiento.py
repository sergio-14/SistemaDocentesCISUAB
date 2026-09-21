from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.page import PageMargins

from .estilos import guardar_en_buffer


AZUL_PLANTILLA = 'FFDBE5F1'
NEGRO = '000000'
BORDE_TABLA = Border(
    left=Side(style='thin', color=NEGRO),
    right=Side(style='thin', color=NEGRO),
    top=Side(style='thin', color=NEGRO),
    bottom=Side(style='thin', color=NEGRO),
)
ANCHOS_PLANTILLA = {
    'A': 30.29, 'B': 33.57, 'C': 27.71, 'D': 21.86, 'E': 3,
    'F': 26, 'G': 21, 'H': 7.14, 'I': 5, 'J': 4.43, 'K': 11.86, 'L': 11.43,
}


def _texto(valor, defecto=''):
    return defecto if valor is None else str(valor).strip()


def _normalizar_fila(fila, indice):
    if not isinstance(fila, dict):
        raise TypeError('Cada fila del seguimiento institucional debe ser un diccionario.')
    normalizada = dict(fila)
    normalizada.setdefault('documento_id', normalizada.get('programa') or 0)
    normalizada.setdefault('indicador_numero', indice)
    normalizada.setdefault('medios_verificacion', [])
    return normalizada


def _normalizar_documentos(documentos):
    if not isinstance(documentos, (list, tuple)):
        raise TypeError('Los documentos del seguimiento institucional deben enviarse como una lista.')
    normalizados = [dict(documento) for documento in documentos]
    return normalizados or [{'id': 0, 'programa': 'Sin programas registrados'}]


def _nombre_hoja(programa, usados):
    prohibidos = set('[]:*?/\\')
    base = ''.join('-' if caracter in prohibidos else caracter for caracter in _texto(programa, 'Seguimiento'))
    base = ' '.join(base.split())[:31] or 'Seguimiento'
    nombre = base
    sufijo = 2
    while nombre.lower() in usados:
        terminacion = f' {sufijo}'
        nombre = f'{base[:31 - len(terminacion)]}{terminacion}'
        sufijo += 1
    usados.add(nombre.lower())
    return nombre


def _aplicar_estilo(celda, *, encabezado=False, centro=False, negrita=False, tamano=11):
    celda.font = Font(name='Calibri', size=tamano, bold=encabezado or negrita, color=NEGRO)
    if encabezado:
        celda.fill = PatternFill('solid', fgColor=AZUL_PLANTILLA)
    celda.border = BORDE_TABLA
    celda.alignment = Alignment(
        horizontal='center' if centro or encabezado else 'left',
        vertical='center' if centro or encabezado else 'top',
        wrap_text=True,
    )


def _agregar_logo(hoja):
    ruta = Path(__file__).resolve().parents[4] / 'frontend' / 'public' / 'images' / 'LOGOUAB.png'
    if not ruta.is_file():
        return
    try:
        imagen = Image(ruta)
        imagen.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=9,
                colOff=pixels_to_EMU(28),
                row=1,
                rowOff=pixels_to_EMU(1),
            ),
            ext=XDRPositiveSize2D(
                cx=pixels_to_EMU(72),
                cy=pixels_to_EMU(65),
            ),
        )
        hoja.add_image(imagen)
    except (OSError, ValueError):
        pass


def _nombre_unidad_sin_facultad(unidad, facultad):
    unidad = _texto(unidad, 'SIN UNIDAD REGISTRADA')
    facultad = _texto(facultad)
    sufijo = f' - {facultad}' if facultad else ''
    if sufijo and unidad.casefold().endswith(sufijo.casefold()):
        unidad = unidad[:-len(sufijo)].strip()
    return unidad


def _configurar_encabezado(hoja, documento, gestion):
    for columna, ancho in ANCHOS_PLANTILLA.items():
        hoja.column_dimensions[columna].width = ancho
    for rango in ('A1:B1', 'A4:K4', 'A5:C5', 'A6:C6', 'A7:C7'):
        hoja.merge_cells(rango)

    hoja['A1'] = 'UNIVERSIDAD AUTÓNOMA DEL BENI JOSÉ BALLIVIÁN'
    hoja['A2'] = 'DIRECCIÓN DE PLANIFICACIÓN Y DESARROLLO ORGANIZACIONAL'
    hoja['A4'] = f'EVALUACIÓN FÍSICA DEL POA POR UNIDAD Y OBJETIVOS ESPECÍFICOS - GESTIÓN {gestion}'
    hoja['A5'] = f"PROGRAMA: {_texto(documento.get('programa'), 'SIN PROGRAMA').upper()}"
    facultad = _texto(documento.get('facultad')).upper()
    unidad = _nombre_unidad_sin_facultad(documento.get('unidad'), facultad).upper()
    partes_ubicacion = [facultad] if facultad else []
    if unidad and unidad.casefold() != facultad.casefold():
        partes_ubicacion.append(unidad)
    ubicacion = ' / '.join(partes_ubicacion)
    hoja['A6'] = f"FACULTAD/CARRERA: {ubicacion or facultad or unidad}"

    for coordenada in ('A1', 'A2'):
        hoja[coordenada].font = Font(name='Calibri', size=12, bold=True)
        hoja[coordenada].alignment = Alignment(horizontal='left', vertical='center')
    hoja['A4'].font = Font(name='Calibri', size=20, bold=True, color=NEGRO)
    hoja['A4'].alignment = Alignment(horizontal='center')
    for coordenada in ('A5', 'A6'):
        hoja[coordenada].font = Font(name='Calibri', size=12, bold=True)
        hoja[coordenada].alignment = Alignment(horizontal='left')
    hoja.row_dimensions[5].height = 21
    hoja.row_dimensions[6].height = 21.75
    hoja.row_dimensions[7].height = 18
    _agregar_logo(hoja)

    encabezados = {
        'A8': 'OBJETIVOS ESPECÍFICOS',
        'B8': 'OPERACIÓN Y/O ACTIVIDAD\n(Inversión / Funcionamiento)',
        'C8': 'PRODUCTOS (Bienes, Normas o Servicios)',
        'C9': '(Resultados logrados)',
        'D9': 'MEDIOS DE VERIFICACIÓN',
        'E8': 'INDICADOR (ACCIÓN / OPERACIÓN)',
        'E9': 'N°',
        'F9': 'DENOMINACIÓN DEL INDICADOR',
        'G9': 'FÓRMULA',
        'H9': f'LÍNEA BASE\n{gestion - 1}',
        'I9': f'META {gestion}',
        'I10': 'Prog.',
        'J10': 'Ejec.',
        'K10': 'Grado de cumplimiento (%)',
    }
    for coordenada, valor in encabezados.items():
        hoja[coordenada] = valor
    for fila in hoja.iter_rows(min_row=8, max_row=10, min_col=1, max_col=11):
        for celda in fila:
            _aplicar_estilo(celda, encabezado=True, centro=True, tamano=11)
    for coordenada in ('I10', 'J10', 'K10'):
        hoja[coordenada].font = Font(name='Calibri', size=10, bold=True, color=NEGRO)
    for rango in (
        'A8:A10', 'B8:B10', 'C8:D8', 'C9:C10', 'D9:D10',
        'E8:K8', 'E9:E10', 'F9:F10', 'G9:G10', 'H9:H10', 'I9:K9',
    ):
        hoja.merge_cells(rango)
    for coordenada in ('E8', 'I9', 'I10'):
        hoja[coordenada].alignment = Alignment(horizontal='center', vertical='center')
    hoja.row_dimensions[8].height = 14.25
    hoja.row_dimensions[9].height = 15
    hoja.row_dimensions[10].height = 41.25


def _alto_fila(valores):
    columnas_texto = ((valores[0], 30), (valores[1], 34), (valores[2], 28), (valores[3], 22), (valores[5], 26))
    lineas = 1
    for valor, ancho in columnas_texto:
        partes = _texto(valor).splitlines() or ['']
        estimadas = sum(max(1, (len(parte) // ancho) + 1) for parte in partes)
        lineas = max(lineas, estimadas)
    return min(180, max(34, lineas * 14))


def _objetivo_texto(fila):
    codigo = _texto(fila.get('objetivo_codigo'))
    descripcion = _texto(fila.get('objetivo_descripcion'), 'Objetivo no especificado')
    return f'{codigo} . {descripcion}' if codigo else descripcion


def _actividad_texto(fila):
    codigo = _texto(fila.get('actividad_codigo'))
    nombre = _texto(fila.get('actividad_nombre'))
    return f'{codigo} . {nombre}' if codigo else nombre


def _medios_texto(fila):
    medios = fila.get('medios_verificacion') or []
    if isinstance(medios, str):
        return medios or 'Sin medios de verificación registrados', [medios] if medios else []
    enlaces = [_texto(medio) for medio in medios if _texto(medio)]
    return ('\n'.join(enlaces), enlaces) if enlaces else ('Sin medios de verificación registrados', [])


def _agregar_imagenes_verificacion(hoja, numero_fila, fila):
    adjuntos = fila.get('medios_archivos') or []
    imagenes = []
    for adjunto in adjuntos:
        ruta = Path(adjunto.get('archivo_local') or '')
        if adjunto.get('tipo') == 'imagen' and ruta.is_file():
            imagenes.append(ruta)
    if not imagenes:
        return 0

    for indice, ruta in enumerate(imagenes[:4]):
        try:
            imagen = Image(ruta)
            factor = min(66 / imagen.width, 58 / imagen.height, 1)
            ancho = max(1, round(imagen.width * factor))
            alto = max(1, round(imagen.height * factor))
            columna_visual = indice % 2
            fila_visual = indice // 2
            imagen.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=3,
                    colOff=pixels_to_EMU(4 + columna_visual * 72),
                    row=numero_fila - 1,
                    rowOff=pixels_to_EMU(4 + fila_visual * 64),
                ),
                ext=XDRPositiveSize2D(
                    cx=pixels_to_EMU(ancho),
                    cy=pixels_to_EMU(alto),
                ),
            )
            hoja.add_image(imagen)
        except (OSError, ValueError):
            continue
    filas_visuales = 1 if len(imagenes) <= 2 else 2
    return (filas_visuales * 64 + 8) * 0.75


def _agregar_datos(hoja, filas):
    fila_inicio = 11
    rangos_objetivos = defaultdict(list)
    if not filas:
        hoja.merge_cells(start_row=fila_inicio, start_column=1, end_row=fila_inicio, end_column=11)
        celda = hoja.cell(fila_inicio, 1, 'No existen actividades registradas para este programa.')
        _aplicar_estilo(celda, centro=True)
        hoja.row_dimensions[fila_inicio].height = 38
        return fila_inicio, fila_inicio

    for indice, fila in enumerate(filas, start=1):
        numero_fila = fila_inicio + indice - 1
        medios, enlaces = _medios_texto(fila)
        valores = [
            _objetivo_texto(fila),
            _actividad_texto(fila),
            _texto(fila.get('resultados_logrados'), 'Sin resultados registrados'),
            medios,
            fila.get('indicador_numero') or indice,
            _texto(fila.get('indicador_descripcion'), 'Sin indicador registrado'),
            'Ejecutado / Programado × 100',
            fila.get('indicador_linea_base') or 0,
            fila.get('programado') or 0,
            fila.get('ejecutado') or 0,
            f'=IFERROR(IF(J{numero_fila}>I{numero_fila},1,J{numero_fila}/I{numero_fila}),0)',
        ]
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(numero_fila, columna, valor)
            _aplicar_estilo(celda)
        alineaciones = {
            1: ('left', 'center', True),
            2: ('left', 'top', True),
            3: ('left', 'center', True),
            4: ('left', 'top', True),
            5: ('center', 'center', True),
            6: (None, 'center', True),
            7: ('center', 'top', True),
            8: ('center', 'center', True),
            9: ('center', 'center', None),
            10: ('center', 'center', True),
            11: ('center', 'center', True),
        }
        for columna, (horizontal, vertical, ajustar) in alineaciones.items():
            hoja.cell(numero_fila, columna).alignment = Alignment(
                horizontal=horizontal,
                vertical=vertical,
                wrap_text=ajustar,
            )
        hoja.cell(numero_fila, 7).number_format = '@'
        hoja.cell(numero_fila, 8).number_format = '@'
        hoja.cell(numero_fila, 11).number_format = '0%'
        hoja.cell(numero_fila, 11).font = Font(name='Arial', size=8, color=NEGRO)
        if len(enlaces) == 1:
            celda_medio = hoja.cell(numero_fila, 4)
            celda_medio.hyperlink = enlaces[0]
            celda_medio.font = Font(name='Calibri', size=11, color='0563C1', underline='single')
        alto_imagenes = _agregar_imagenes_verificacion(hoja, numero_fila, fila)
        hoja.row_dimensions[numero_fila].height = max(_alto_fila(valores), alto_imagenes)
        rangos_objetivos[fila.get('objetivo_id') or _objetivo_texto(fila)].append(numero_fila)

    for numeros in rangos_objetivos.values():
        if len(numeros) > 1 and numeros == list(range(numeros[0], numeros[-1] + 1)):
            hoja.merge_cells(start_row=numeros[0], start_column=1, end_row=numeros[-1], end_column=1)
            hoja.cell(numeros[0], 1).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    return fila_inicio, fila_inicio + len(filas) - 1


def _agregar_total_y_firmas(hoja, documento, fila_inicio, fila_fin):
    fila_total = fila_fin + 1
    hoja.merge_cells(start_row=fila_total, start_column=1, end_row=fila_total, end_column=10)
    hoja.cell(fila_total, 1, 'TOTAL GENERAL')
    hoja.cell(fila_total, 11, f'=IFERROR(AVERAGE(K{fila_inicio}:K{fila_fin}),0)')
    for celda in hoja[fila_total][:11]:
        _aplicar_estilo(celda, encabezado=True, centro=True)
    hoja.cell(fila_total, 11).number_format = '0%'

    fila_firma = fila_total + 2
    hoja.merge_cells(start_row=fila_firma, start_column=1, end_row=fila_firma, end_column=6)
    hoja.merge_cells(start_row=fila_firma, start_column=7, end_row=fila_firma, end_column=11)
    hoja.cell(fila_firma, 7, 'FIRMA')
    fecha = documento.get('fecha_elaboracion') or date.today()
    datos_firma = (
        ('RESPONSABLE DE UNIDAD', documento.get('jefe_unidad') or ''),
        ('PROFESIONAL QUE ELABORA', documento.get('elaborado_por') or ''),
        ('FECHA DE ELABORACIÓN', fecha),
    )
    for desplazamiento, (etiqueta, valor) in enumerate(datos_firma, start=1):
        numero_fila = fila_firma + desplazamiento
        hoja.cell(numero_fila, 1, etiqueta)
        hoja.merge_cells(start_row=numero_fila, start_column=2, end_row=numero_fila, end_column=6)
        hoja.cell(numero_fila, 2, valor)
        hoja.merge_cells(start_row=numero_fila, start_column=7, end_row=numero_fila, end_column=11)
    for fila in hoja.iter_rows(min_row=fila_firma, max_row=fila_firma + 3, min_col=1, max_col=11):
        for celda in fila:
            _aplicar_estilo(celda, centro=celda.column >= 2, negrita=True)
    for celda in hoja[fila_firma][:11]:
        celda.fill = PatternFill('solid', fgColor=AZUL_PLANTILLA)
        celda.font = Font(name='Calibri', size=10, bold=True, color=NEGRO)
    hoja.row_dimensions[fila_firma].height = 15.75
    hoja.row_dimensions[fila_firma + 1].height = 45
    hoja.row_dimensions[fila_firma + 2].height = 45
    hoja.row_dimensions[fila_firma + 3].height = 15.75
    return fila_firma + 3


def _configurar_impresion(hoja):
    hoja.freeze_panes = None
    hoja.sheet_view.showGridLines = True
    hoja.sheet_view.zoomScale = 100
    hoja.page_setup.orientation = 'landscape'
    hoja.page_setup.paperSize = 14
    hoja.page_margins = PageMargins(
        left=0.5905511811023623,
        right=0.1968503937007874,
        top=0.9448818897637796,
        bottom=0.9448818897637796,
        header=0,
        footer=0,
    )


def _crear_hoja_programa(libro, documento, filas, gestion, usados):
    hoja = libro.create_sheet(_nombre_hoja(documento.get('programa'), usados))
    _configurar_encabezado(hoja, documento, gestion)
    fila_inicio, fila_fin = _agregar_datos(hoja, filas)
    _agregar_total_y_firmas(hoja, documento, fila_inicio, fila_fin)
    _configurar_impresion(hoja)


def generar_seguimiento_institucional_excel(filas, gestion, documentos):
    filas = [_normalizar_fila(fila, indice) for indice, fila in enumerate(filas, start=1)]
    documentos = _normalizar_documentos(documentos)
    filas_por_documento = defaultdict(list)
    for fila in filas:
        filas_por_documento[fila.get('documento_id')].append(fila)

    libro = Workbook()
    libro.remove(libro.active)
    usados = set()
    for documento in documentos:
        _crear_hoja_programa(
            libro, documento, filas_por_documento.get(documento.get('id'), []), gestion, usados,
        )
    libro.active = 0
    libro.calculation.fullCalcOnLoad = True
    libro.calculation.forceFullCalc = True
    libro.calculation.calcMode = 'auto'
    return guardar_en_buffer(libro)
