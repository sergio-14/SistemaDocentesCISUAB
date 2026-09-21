from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from django.contrib.auth.models import User
from django.db import IntegrityError, transaction
from django.test import RequestFactory, SimpleTestCase, TestCase
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory

from fondos.models import Carrera
from poa_document.api.serializers import ActividadSerializer, DetallePresupuestoSerializer
from poa_document.api.views import (
    TableroPOAView,
    _consolidar_requerimientos,
    _errores_formulacion_documento,
    _respuesta_pdf,
)
from poa_document.models import Actividad, DetallePresupuesto, DocumentoPOA, ObjetivoEspecifico, ProgramaPOA, SeguimientoActividadPOA, VersionDocumentoPOA
from poa_document.reportes.excel import (
    generar_catalogo_items_excel,
    generar_consolidado_requerimientos_excel,
    generar_seguimiento_institucional_excel,
)
from poa_document.reportes.pdf import (
    DocumentoPOAPDFGenerator,
    generar_consolidado_requerimientos_pdf,
    generar_seguimiento_institucional_pdf,
)


class ProgramaPOAModelTests(TestCase):
    def setUp(self):
        self.carrera_a = Carrera.objects.create(
            nombre='Carrera POA A', codigo='PAA', facultad='Facultad de prueba A'
        )
        self.carrera_b = Carrera.objects.create(
            nombre='Carrera POA B', codigo='PAB', facultad='Facultad de prueba B'
        )

    def test_el_nombre_de_programa_es_unico_dentro_de_una_carrera(self):
        ProgramaPOA.objects.create(carrera=self.carrera_a, nombre='Programa Regular')

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                ProgramaPOA.objects.create(carrera=self.carrera_a, nombre='Programa Regular')

    def test_el_mismo_nombre_se_permite_en_otra_carrera(self):
        ProgramaPOA.objects.create(carrera=self.carrera_a, nombre='Programa Regular')
        programa = ProgramaPOA.objects.create(carrera=self.carrera_b, nombre='Programa Regular')

        self.assertEqual(programa.carrera_id, self.carrera_b.id)


class FormulacionPOATests(TestCase):
    def setUp(self):
        self.carrera = Carrera.objects.create(nombre='Carrera Calidad', codigo='PQC', facultad='Prueba')
        self.documento = DocumentoPOA.objects.create(
            gestion=2030, unidad_solicitante=self.carrera, programa='Programa de prueba',
            objetivo_gestion_institucional='Mejorar la gestión', fecha_elaboracion=date(2030, 1, 1),
        )
        self.objetivo = ObjetivoEspecifico.objects.create(documento=self.documento, codigo='OE-1', descripcion='Objetivo')

    def test_no_permite_mes_final_anterior_al_inicio(self):
        serializer = ActividadSerializer(data={
            'objetivo_id': self.objetivo.id, 'codigo': 'A-1', 'nombre': 'Actividad', 'responsable': 'Responsable',
            'productos_esperados': 'Producto', 'mes_inicio': 'octubre', 'mes_fin': 'marzo',
            'indicador_linea_base': 0, 'indicador_meta': 1,
        })
        self.assertFalse(serializer.is_valid())
        self.assertIn('mes_fin', serializer.errors)

    def test_revision_detecta_indicador_obligatorio(self):
        Actividad.objects.create(
            objetivo=self.objetivo, codigo='A-1', nombre='Actividad', responsable='Responsable',
            productos_esperados='Producto', mes_inicio='enero', mes_fin='febrero',
            indicador_descripcion='', indicador_linea_base=0, indicador_meta=1,
        )
        errores = _errores_formulacion_documento(self.documento)
        self.assertTrue(any('debe registrar indicador' in error for error in errores))

    def test_presupuesto_exige_cantidad_positiva(self):
        serializer = DetallePresupuestoSerializer(data={'cantidad': 0, 'costo_unitario': 10})
        self.assertFalse(serializer.is_valid())
        self.assertIn('cantidad', serializer.errors)

    def test_version_y_seguimiento_conservan_trazabilidad(self):
        user = User.objects.create_user(username='director-prueba')
        version = VersionDocumentoPOA.objects.create(
            documento=self.documento, numero=1, motivo='Aprobación', creado_por=user, snapshot={'objetivos': []},
        )
        actividad = Actividad.objects.create(
            objetivo=self.objetivo, codigo='A-2', nombre='Actividad', responsable='Responsable',
            productos_esperados='Producto', mes_inicio='enero', mes_fin='febrero',
            indicador_linea_base=0, indicador_meta=1,
        )
        seguimiento = SeguimientoActividadPOA.objects.create(
            actividad=actividad, estado_anterior='programado', estado_nuevo='en_ejecucion', registrado_por=user,
        )
        self.assertTrue(version.vigente)
        self.assertEqual(seguimiento.actividad_id, actividad.id)

    def test_consolidado_separa_partidas_y_suma_item_igual(self):
        self.documento.estado = 'aprobado'; self.documento.save(update_fields=['estado'])
        actividad = Actividad.objects.create(
            objetivo=self.objetivo, codigo='A-3', nombre='Compra', responsable='Resp', productos_esperados='P',
            mes_inicio='enero', mes_fin='febrero', indicador_linea_base=0, indicador_meta=1,
        )
        DetallePresupuesto.objects.create(actividad=actividad, tipo='funcionamiento', partida='25100', item='Papel', unidad_medida='Paquete', cantidad=2, costo_unitario=10, costo_total=0, mes_requerimiento='enero')
        DetallePresupuesto.objects.create(actividad=actividad, tipo='funcionamiento', partida='25100', item='Papel', unidad_medida='Paquete', cantidad=3, costo_unitario=10, costo_total=0, mes_requerimiento='febrero')
        request = APIRequestFactory().get('/api/poa/consolidado-requerimientos/', {'gestion': 2030})
        request.user = User.objects.create_user(username='admin-consolidado', is_superuser=True)
        _, items = _consolidar_requerimientos(request)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]['cantidad_total'], 5)

    def test_tablero_respeta_gestion_y_expone_presupuesto(self):
        self.documento.estado = 'ejecucion'; self.documento.save(update_fields=['estado'])
        request = APIRequestFactory().get('/api/poa/tablero/', {'gestion': 2030})
        request.user = User.objects.create_user(username='admin-tablero', is_superuser=True)
        response = TableroPOAView().get(request)
        self.assertEqual(response.status_code, 200)
        self.assertIn('presupuesto', response.data)


class GeneradoresReportesPOATests(SimpleTestCase):
    items = [{
        'partida': '100',
        'tipo': 'funcion',
        'item': 'Material A&B <especial>',
        'unidad_medida': 'Unidad',
        'caracteristicas': 'Prueba',
        'cantidad_total': 2,
        'monto_estimado_total': Decimal('25.50'),
        'actividades_solicitantes': 1,
    }]
    filas_seguimiento_pdf = [[
        'Programa', 'A-1', 'Actividad <prueba>', 'Responsable', 'Enero', 'Febrero',
        'completado', 75.0, '100.00', 'Resultado A&B', 10, 8,
    ]]
    filas_seguimiento = [{
        'documento_id': 1,
        'programa': '100 Formación Pregrado',
        'unidad': 'Ingeniería de Sistemas - Facultad de Ingeniería',
        'facultad': 'Facultad de Ingeniería',
        'objetivo_id': 1,
        'objetivo_codigo': 'OE 1',
        'objetivo_descripcion': 'Mejorar la formación profesional',
        'actividad_codigo': '100-0-1',
        'actividad_nombre': 'Actividad <prueba>',
        'resultados_logrados': 'Resultado A&B',
        'medios_verificacion': ['https://example.test/a'],
        'indicador_numero': 1,
        'indicador_descripcion': 'Número de actividades completadas',
        'indicador_linea_base': 2,
        'programado': 10,
        'ejecutado': 8,
        'grado_cumplimiento': 80,
        'presupuesto': '100.00',
    }]
    documentos_seguimiento = [{
        'id': 1,
        'programa': '100 Formación Pregrado',
        'unidad': 'Ingeniería de Sistemas',
        'facultad': 'Facultad de Ingeniería',
        'jefe_unidad': 'Responsable de unidad',
        'elaborado_por': 'Profesional POA',
        'fecha_elaboracion': '2026-08-30',
    }]
    def test_respuesta_pdf_respeta_el_modo_del_visualizador(self):
        request_factory = RequestFactory()
        vista_previa = _respuesta_pdf(
            request_factory.get('/reporte/?view=true'), BytesIO(b'%PDF'), 'reporte.pdf',
        )
        descarga_compatible = _respuesta_pdf(
            request_factory.get('/reporte/'), BytesIO(b'%PDF'), 'reporte.pdf',
        )
        self.assertTrue(vista_previa['Content-Disposition'].startswith('inline;'))
        self.assertTrue(descarga_compatible['Content-Disposition'].startswith('attachment;'))

    def test_pdfs_generados_son_archivos_validos(self):
        archivos = [
            DocumentoPOAPDFGenerator.generar_reporte_general([], 2026),
            generar_consolidado_requerimientos_pdf(self.items, 2026),
            generar_seguimiento_institucional_pdf(self.filas_seguimiento_pdf, 2026),
        ]
        for archivo in archivos:
            self.assertEqual(archivo.read(4), b'%PDF')

    def test_excel_consolidado_conserva_hoja_y_total(self):
        archivo = generar_consolidado_requerimientos_excel(self.items, 2026)
        libro = load_workbook(archivo, data_only=True)
        self.assertEqual(libro.sheetnames, ['Consolidado compras'])
        self.assertEqual(libro.active.cell(libro.active.max_row, 7).value, 25.5)

    def test_excel_seguimiento_copia_la_matriz_institucional(self):
        filas = [
            self.filas_seguimiento[0],
            {
                **self.filas_seguimiento[0],
                'actividad_codigo': '100-0-2',
                'actividad_nombre': 'Segunda actividad',
                'indicador_numero': 2,
            },
            {
                **self.filas_seguimiento[0],
                'documento_id': 2,
                'programa': '510 Formación Investigación',
                'objetivo_id': 2,
                'actividad_codigo': '510-0-1',
                'actividad_nombre': 'Actividad de investigación',
            },
        ]
        documentos = [
            self.documentos_seguimiento[0],
            {
                **self.documentos_seguimiento[0],
                'id': 2,
                'programa': '510 Formación Investigación',
            },
        ]
        archivo = generar_seguimiento_institucional_excel(
            filas, 2026, documentos,
        )
        libro = load_workbook(archivo, data_only=False)
        self.assertEqual(
            libro.sheetnames,
            ['100 Formación Pregrado', '510 Formación Investigación'],
        )
        hoja = libro['100 Formación Pregrado']
        self.assertEqual(hoja['A1'].value, 'UNIVERSIDAD AUTÓNOMA DEL BENI JOSÉ BALLIVIÁN')
        self.assertEqual(hoja['A8'].value, 'OBJETIVOS ESPECÍFICOS')
        self.assertEqual(hoja['E8'].value, 'INDICADOR (ACCIÓN / OPERACIÓN)')
        self.assertEqual(
            hoja['A6'].value,
            'FACULTAD/CARRERA: FACULTAD DE INGENIERÍA / INGENIERÍA DE SISTEMAS',
        )
        self.assertEqual(hoja['A11'].value, 'OE 1 . Mejorar la formación profesional')
        self.assertEqual(hoja['B11'].value, '100-0-1 . Actividad <prueba>')
        self.assertEqual(hoja['K11'].value, '=IFERROR(IF(J11>I11,1,J11/I11),0)')
        self.assertIsNone(hoja.freeze_panes)
        self.assertEqual(hoja.page_setup.orientation, 'landscape')
        self.assertEqual(hoja.page_setup.paperSize, 14)
        self.assertEqual(hoja['A4'].font.name, 'Calibri')
        self.assertEqual(hoja['A4'].font.sz, 20)
        self.assertTrue(hoja['A4'].font.bold)
        self.assertIsNone(hoja['A4'].fill.patternType)
        self.assertEqual(hoja['A8'].font.sz, 11)
        self.assertEqual(hoja['A8'].fill.fgColor.rgb, 'FFDBE5F1')
        self.assertEqual(hoja.column_dimensions['E'].width, 3)
        self.assertEqual(hoja.column_dimensions['J'].width, 4.43)
        self.assertEqual(hoja.row_dimensions[8].height, 14.25)
        self.assertEqual(hoja.row_dimensions[9].height, 15)
        self.assertEqual(hoja.row_dimensions[10].height, 41.25)
        self.assertTrue(hoja.sheet_view.showGridLines)
        self.assertEqual(len(hoja._images), 1)
        self.assertIn('A1:B1', {str(rango) for rango in hoja.merged_cells.ranges})
        self.assertIn('A8:A10', {str(rango) for rango in hoja.merged_cells.ranges})
        self.assertIn('A11:A12', {str(rango) for rango in hoja.merged_cells.ranges})

    def test_excel_catalogo_mantiene_columnas_publicas(self):
        archivo = generar_catalogo_items_excel([
            SimpleNamespace(detalle='Papel', partida='200', unidad_medida='Paquete'),
        ])
        libro = load_workbook(archivo, data_only=True)
        self.assertEqual(
            [celda.value for celda in libro.active[1]],
            ['DETALLE', 'partida', 'UNIDAD_MEDIDA'],
        )
